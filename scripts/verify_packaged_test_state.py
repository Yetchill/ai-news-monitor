#!/usr/bin/env python3
"""Verify durable packaged-test database state and reopen Office exports."""

from __future__ import annotations

import argparse
import json
import sqlite3
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

MARKER = "打包集成持久化标记-PACKAGED_EXPORT_MARKER"
ALPHA = "PACKAGED_CRAWLER_ALPHA"
BETA = "PACKAGED_CRAWLER_BETA"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database", required=True, type=Path)
    parser.add_argument("--phase", choices=("integration", "restarted"), required=True)
    parser.add_argument("--xlsx", type=Path)
    parser.add_argument("--docx", type=Path)
    args = parser.parse_args()
    if (args.xlsx is None) != (args.docx is None):
        parser.error("--xlsx and --docx must be supplied together")

    database_path = args.database.resolve()
    if not database_path.is_file():
        raise AssertionError(f"database does not exist: {database_path}")
    summary = _verify_database(database_path, phase=args.phase)
    if args.xlsx is not None and args.docx is not None:
        _verify_xlsx(args.xlsx.resolve())
        _verify_docx(args.docx.resolve())
        summary["office_exports_reopened"] = True
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    return 0


def _verify_database(path: Path, *, phase: str) -> dict[str, object]:
    with sqlite3.connect(path, timeout=5) as connection:
        connection.row_factory = sqlite3.Row
        runs = connection.execute(
            "SELECT status, new_count, duplicate_count, finished_at "
            "FROM crawl_runs WHERE trigger = 'scheduled' ORDER BY id"
        ).fetchall()
        active_runs = sum(row["status"] == "running" or row["finished_at"] is None for row in runs)
        failed_runs = sum(row["status"] in {"failed", "partial_success"} for row in runs)
        successful_runs = sum(row["status"] == "success" for row in runs)
        if len(runs) < 4:
            raise AssertionError(f"expected at least four scheduled runs, found {len(runs)}")
        if active_runs:
            raise AssertionError(f"found {active_runs} stuck scheduled CrawlRun rows")
        if failed_runs < 1 or successful_runs < 2:
            raise AssertionError(
                "expected failure recovery, "
                f"failed/partial={failed_runs}, success={successful_runs}"
            )
        if not any(row["new_count"] > 0 for row in runs):
            raise AssertionError("scheduled crawler never persisted a new item")
        if not any(row["duplicate_count"] > 0 for row in runs):
            raise AssertionError("scheduled crawler never recorded a duplicate")

        item_rows = connection.execute(
            "SELECT title, ai_summary FROM intelligence_items WHERE is_active = 1"
        ).fetchall()
        titles = [str(row["title"]) for row in item_rows]
        for expected in (MARKER, ALPHA, BETA):
            if not any(expected in title for title in titles):
                raise AssertionError(f"missing persisted packaged item: {expected}")
        if not any(
            row["ai_summary"] and "零外部 Token" in str(row["ai_summary"]) for row in item_rows
        ):
            raise AssertionError("fake packaged AI summary was not persisted")

        source = connection.execute(
            "SELECT source_kind, origin, enabled, last_success_at, last_error "
            "FROM sources WHERE slug = 'packaged-integration-fixture'"
        ).fetchone()
        if source is None:
            raise AssertionError("formal packaged fixture source is missing")
        if (
            source["source_kind"] != "formal"
            or source["origin"] != "preset"
            or not source["enabled"]
            or source["last_success_at"] is None
            or source["last_error"] is not None
        ):
            raise AssertionError(f"fixture source did not recover durably: {dict(source)}")

        schedule = connection.execute(
            "SELECT schedule_enabled, last_scheduled_trigger_at FROM schedule_settings WHERE id = 1"
        ).fetchone()
        if schedule is None or not schedule["schedule_enabled"]:
            raise AssertionError("persisted schedule is not enabled")
        if schedule["last_scheduled_trigger_at"] is None:
            raise AssertionError("scheduler did not persist its last trigger target")

        active_ai = connection.execute(
            "SELECT COUNT(*) FROM ai_jobs WHERE status IN ('pending', 'running')"
        ).fetchone()[0]
        completed_ai = connection.execute(
            "SELECT COUNT(*) FROM ai_jobs "
            "WHERE job_type = 'summarization' AND status = 'completed' "
            "AND success_count > 0 AND provider = 'custom' "
            "AND model = 'packaged-zero-token-fixture'"
        ).fetchone()[0]
        if active_ai:
            raise AssertionError(f"found {active_ai} stuck packaged AI jobs")
        if completed_ai < 1:
            raise AssertionError("no completed packaged fake-AI summarization job was found")

    return {
        "phase": phase,
        "scheduled_runs": len(runs),
        "failed_or_partial_runs": failed_runs,
        "successful_runs": successful_runs,
        "new_items": sum(int(row["new_count"]) for row in runs),
        "duplicates": sum(int(row["duplicate_count"]) for row in runs),
        "active_runs": active_runs,
        "active_ai_jobs": active_ai,
        "completed_fake_ai_jobs": completed_ai,
    }


def _verify_xlsx(path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = (
            str(cell)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
            for cell in row
            if cell is not None
        )
        if not any(MARKER in value for value in values):
            raise AssertionError("reopened XLSX does not contain the Chinese marker")
    finally:
        workbook.close()


def _verify_docx(path: Path) -> None:
    document = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if MARKER not in text:
        raise AssertionError("reopened DOCX does not contain the Chinese marker")


if __name__ == "__main__":
    raise SystemExit(main())
